# -*- coding: utf-8 -*-
"""通用数据表（报销明细表）。

这是两层架构的**契约**：上游不管材料来自哪个渠道、什么费用类型，
核对完都落成这一张表；下游任何专用模版都从这张表取数。

它同时是给人看的**核对底稿**——一行一笔、嵌完整截图、标明发票来源与差额，
财务据此复核，而不必去读特定格式的台账。
"""
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .ledger import make_thumb

F = "微软雅黑"
_thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
HDR_FONT = Font(name=F, bold=True, color="FFFFFF", size=11)
HDR_FILL = PatternFill("solid", start_color="305496")
BODY = Font(name=F, size=11)
BOLD = Font(name=F, bold=True, size=11)
DIFF = Font(name=F, bold=True, size=11, color="C55A11")
RED = Font(name=F, bold=True, size=11, color="C00000")
NOINV = PatternFill("solid", start_color="FFF2CC")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

# 通用数据表的字段。专用模版的映射关系以此为准。
FIELDS = ["序号", "消费时间", "实付金额(元)", "支付截图", "费用具体明细",
          "费用类别", "发票来源", "发票金额(元)", "差额(元)", "凭证类型"]
WIDTHS = [6, 19, 13, 36, 40, 11, 28, 13, 11, 12]


def _header(ws, cols, widths):
    from openpyxl.utils import get_column_letter
    for c, t in enumerate(cols, 1):
        cell = ws.cell(row=1, column=c, value=t)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.alignment = CENTER
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = "A2"


def build(rows, images, summary, out_path, cache_dir=None):
    """rows 为条目列表，images 为 {截图文件名: jpeg字节}。

    summary 取自 manifest.settle()，键名以那里为准：
    total / n_total / inv_ledger / sub_total / n_inv / n_sub /
    actual_inv / missing_inv。
    """
    wb = Workbook()

    # ---- 报销明细 ----
    ws = wb.active
    ws.title = "报销明细"
    _header(ws, FIELDS, WIDTHS)
    for i, r in enumerate(rows, 1):
        n = i + 1
        diff = r.get("diff")
        vals = [i, _fmt_date(r.get("date")), r.get("amount"), None, r.get("detail"),
                r.get("category"), r.get("invoice") or "无（需替票）",
                r.get("invoice_amount"), diff or None, r.get("voucher")]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=n, column=c, value=v)
            cell.font = DIFF if (c == 9 and diff) else BODY
            cell.border = BORDER
            if c != 4:
                cell.alignment = CENTER
            if c in (3, 8):
                cell.number_format = "#,##0.00"
            if c == 9:
                cell.number_format = "+#,##0.00;-#,##0.00"
        if not r.get("invoice"):
            for c in range(1, len(FIELDS) + 1):
                ws.cell(row=n, column=c).fill = NOINV
        shot = r.get("shot")
        if shot and shot in images and cache_dir:
            path = _write_tmp(images[shot], cache_dir, shot)
            img = XLImage(path)
            img.width, img.height = r.get("img_w", 300), r.get("img_h", 536)
            ws.add_image(img, f"D{n}")
            ws.row_dimensions[n].height = img.height * 0.75 + 4
        else:
            ws.row_dimensions[n].height = 22
    last = len(rows) + 1
    tr = last + 1
    ws.cell(row=tr, column=2, value="合计").font = BOLD
    for col in ("C", "H", "I"):
        cell = ws.cell(row=tr, column=ord(col) - 64, value=f"=SUM({col}2:{col}{last})")
        cell.font = BOLD
        cell.number_format = "#,##0.00"
    for c in range(1, len(FIELDS) + 1):
        ws.cell(row=tr, column=c).border = BORDER
        ws.cell(row=tr, column=c).alignment = CENTER

    _summary_sheet(wb, rows, summary)
    _notes_sheet(wb, rows, summary)
    wb.save(out_path)
    return out_path


def _fmt_date(v):
    return v.strftime("%Y-%m-%d") if hasattr(v, "strftime") else (v or "")


def _write_tmp(blob, cache_dir, name):
    import os
    os.makedirs(cache_dir, exist_ok=True)
    p = os.path.join(cache_dir, name.rsplit(".", 1)[0] + "_ds.jpg")
    with open(p, "wb") as fh:
        fh.write(blob)
    return p


def _summary_sheet(wb, rows, s):
    ws = wb.create_sheet("金额汇总")
    _header(ws, ["维度", "项目", "笔数", "金额(元)", "说明"],
            [14, 30, 8, 16, 40])
    n = 1

    def put(dim, item, cnt, amt, note=""):
        nonlocal n
        n += 1
        for c, v in enumerate([dim, item, cnt, amt, note], 1):
            cell = ws.cell(row=n, column=c, value=v)
            cell.font = BODY
            cell.border = BORDER
            cell.alignment = CENTER
            if c == 4:
                cell.number_format = "#,##0.00"
        ws.row_dimensions[n].height = 22

    put("总额", "报销总额", s["n_total"], s["total"], "实付合计，报销依据")
    put("凭证", "有发票部分", s.get("n_inv", 0), s["inv_ledger"],
        "可凭发票抵扣")
    put("凭证", "需替票部分", s.get("n_sub", 0), s["sub_total"],
        "无发票或发票不足额的部分")
    by_cat = defaultdict(lambda: [0, 0.0])
    for r in rows:
        g = by_cat[r.get("category") or "未分类"]
        g[0] += 1
        g[1] = round(g[1] + (r.get("amount") or 0), 2)
    for cat, (cnt, amt) in sorted(by_cat.items(), key=lambda x: -x[1][1]):
        put("费用类别", cat, cnt, amt)
    put("发票", "实际发票金额", "", s.get("actual_inv", 0), "发票 PDF 合计")
    put("发票", "缺少发票金额", "", max(s.get("missing_inv", 0), 0),
        "台账应开减去实际发票")


def _notes_sheet(wb, rows, s):
    ws = wb.create_sheet("核对说明")
    _header(ws, ["消费时间", "实付金额(元)", "发票金额(元)", "差额(元)", "说明"],
            [19, 14, 14, 12, 50])
    n = 1
    for r in rows:
        d = r.get("diff")
        if not d:
            continue
        n += 1
        why = "无发票，需全额替票" if not r.get("invoice") else "实付高于发票额，差额需替票"
        for c, v in enumerate([_fmt_date(r.get("date")), r.get("amount"),
                               r.get("invoice_amount"), d, why], 1):
            cell = ws.cell(row=n, column=c, value=v)
            cell.font = BODY
            cell.border = BORDER
            cell.alignment = CENTER
            if c in (2, 3):
                cell.number_format = "#,##0.00"
            if c == 4:
                cell.number_format = "+#,##0.00;-#,##0.00"
        ws.row_dimensions[n].height = 22
    for j, txt in enumerate([
        "金额口径：报销金额=支付截图实付金额；发票金额仅为可抵扣部分，两者差额需替票。",
        f"合计 {s['n_total']} 笔 {s['total']:.2f} 元，其中有发票 "
        f"{s['inv_ledger']:.2f} 元、需替票 {s['sub_total']:.2f} 元。",
        "本表为核对底稿，专用台账由此表取数生成；两者金额应完全一致。",
    ]):
        cell = ws.cell(row=n + 3 + j, column=1, value=txt)
        cell.font = RED
        ws.merge_cells(start_row=n + 3 + j, start_column=1,
                       end_row=n + 3 + j, end_column=5)
