# -*- coding: utf-8 -*-
"""按 spec 渲染专用台账。运行时纯代码，不调用模型。"""
import logging

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter

from .ledger import inject_cellimages, make_thumb
from .template_spec import TemplateSpec

log = logging.getLogger(__name__)

F = "宋体"
_thin = Side(style="thin", color="000000")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _value(col, row, ctx):
    """按列规则取值。三种方式互斥，spec 校验已保证。"""
    if col.const is not None:
        return col.const
    if col.source:
        return row.get(col.source, ctx.get(col.source))
    if col.template:
        if not _cond(col.when, row):
            return None
        try:
            return col.template.format(**row)
        except (KeyError, ValueError, IndexError):
            return None
    return None


def _cond(when, row):
    """条件判断。默认恒真——没写条件就一直填。"""
    if not when or when == "always":
        return True
    diff = row.get("diff") or 0
    if when == "diff_positive":
        return diff > 0
    if when == "no_invoice":
        return not row.get("invoice")
    if when == "has_invoice":
        return bool(row.get("invoice"))
    return True


def render(spec: TemplateSpec, rows, ctx, out_path, images=None, cache_dir=None):
    """rows 为通用数据表的行，ctx 提供需求人等全局字段。"""
    wb = Workbook()
    ws = wb.active
    ws.title = spec.sheet_name

    for h in spec.header_rows:
        r = h.get("row", 1)
        for cell_def in h.get("cells", []):
            c = ws.cell(row=r, column=column_index_from_string(cell_def["col"]),
                        value=cell_def.get("text"))
            c.font = Font(name=F, bold=True, size=cell_def.get("size", 10))
            c.alignment = CENTER
            c.border = BORDER
            if cell_def.get("fill"):
                c.fill = PatternFill("solid", start_color=cell_def["fill"])
        for m in h.get("merges", []):
            ws.merge_cells(m)
        if h.get("height"):
            ws.row_dimensions[r].height = h["height"]

    for col in spec.columns:
        if col.width:
            ws.column_dimensions[col.letter].width = col.width

    body = Font(name=F, size=10)
    start = spec.data_start_row
    ncol = max((column_index_from_string(c.letter) for c in spec.columns), default=1)
    for i, row in enumerate(rows):
        n = start + i
        for col in spec.columns:
            idx = column_index_from_string(col.letter)
            cell = ws.cell(row=n, column=idx, value=_value(col, row, ctx))
            cell.font = body
            cell.border = BORDER
            cell.alignment = LEFT if col.source == "detail" else CENTER
            if col.number_format:
                cell.number_format = col.number_format
        for c in range(1, ncol + 1):
            ws.cell(row=n, column=c).border = BORDER
        ws.row_dimensions[n].height = 20

    last = start + len(rows) - 1
    if spec.total_row and rows:
        tr = last + 1
        t = spec.total_row
        lc = column_index_from_string(t.get("label_col", "A"))
        ws.cell(row=tr, column=lc, value=t.get("label", "合计")).font = \
            Font(name=F, bold=True, size=10)
        for cl in t.get("sum_cols", []):
            cell = ws.cell(row=tr, column=column_index_from_string(cl),
                           value=f"=SUM({cl}{start}:{cl}{last})")
            cell.font = Font(name=F, bold=True, size=10)
            cell.number_format = "#,##0.00"
        for m in t.get("merges", []):
            ws.merge_cells(m.format(row=tr))
        for c in range(1, ncol + 1):
            ws.cell(row=tr, column=c).border = BORDER
            ws.cell(row=tr, column=c).alignment = CENTER

    ids = _image_sheet(wb, spec, rows, images or {}, cache_dir)
    for name in spec.extra_sheets:
        wb.create_sheet(name)

    wb.save(out_path)
    if ids:
        inject_cellimages(out_path, ids)      # WPS DISPIMG 需在 zip 层注入
    return out_path


def _image_sheet(wb, spec, rows, images, cache_dir):
    """生成截图页。WPS 用 DISPIMG 单元格图片，图片行与金额行交替。"""
    cfg = spec.image_sheet
    if not cfg or not images:
        return []
    import uuid
    ws = wb.create_sheet(cfg.get("name", "付款截图"))
    per = int(cfg.get("per_row", 20))
    for c in range(1, per + 1):
        ws.column_dimensions[get_column_letter(c)].width = 18.775
    ids = []
    for i, row in enumerate(rows):
        band, col = divmod(i, per)
        img_row, amt_row = band * 2 + 1, band * 2 + 2
        shot = row.get("shot")
        if shot and shot in images:
            iid = "ID_" + uuid.uuid4().hex.upper()
            ids.append((iid, images[shot]))
            ws.cell(row=img_row, column=col + 1,
                    value=f'=_xlfn.DISPIMG("{iid}",1)')
        a = ws.cell(row=amt_row, column=col + 1, value=row.get("amount"))
        a.number_format = "#,##0.00"
        a.alignment = CENTER
        a.font = Font(name=F, size=10)
        ws.row_dimensions[img_row].height = 231
        ws.row_dimensions[amt_row].height = 18
    return ids
